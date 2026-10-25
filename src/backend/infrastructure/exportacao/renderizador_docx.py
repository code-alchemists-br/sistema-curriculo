"""Gera o arquivo DOCX do currículo a partir do documento renderizado.

O módulo implementa a porta ``RenderizadorDocumentoCurriculo`` com a biblioteca
python-docx, desenhando cabeçalho, seções e blocos na ordem recebida, com
estilos de título navegáveis e variação de tamanho entre os modelos clássico e
compacto. Ele existe para atender ao UC08 e ao CT-RF09-02 sem que a Application
conheça a biblioteca de geração de Word.
"""

from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.document import Document as DocumentoWord
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from backend.application.documento_renderizado import (
    BlocoDocumento,
    CabecalhoDocumento,
    DocumentoCurriculoRenderizado,
)

IDIOMA_DOCUMENTO = "pt-BR"
FONTE_DOCUMENTO = "Calibri"
SEPARADOR_CONTATOS = " | "


@dataclass(frozen=True, slots=True)
class EstiloModelo:
    """Agrupa as medidas visuais de um modelo de currículo no DOCX.

    O DTO guarda tamanho de fonte, espaçamento entre parágrafos e margem da
    página. Ele existe para que os modelos mudem apenas a aparência, mantendo
    exatamente o mesmo texto.
    """

    tamanho_fonte: float
    espaco_apos: float
    margem_cm: float


ESTILO_CLASSICO = EstiloModelo(tamanho_fonte=11, espaco_apos=6, margem_cm=2.0)
ESTILO_COMPACTO = EstiloModelo(tamanho_fonte=9.5, espaco_apos=2, margem_cm=1.5)
ESTILOS_POR_MODELO = {"compacto": ESTILO_COMPACTO}


class RenderizadorDocumentoCurriculoDocx:
    """Desenha o documento renderizado do currículo em um arquivo Word (DOCX).

    O renderizador cria um documento novo, aplica fonte, idioma e margens do
    modelo, escreve o nome como título, os contatos em uma linha e cada seção
    como "Título 1", seguida dos blocos com título em negrito, subtítulo,
    período e detalhes; ele não decide conteúdo. Ele existe para implementar a
    porta de renderização no formato DOCX.
    """

    def renderizar(self, documento: DocumentoCurriculoRenderizado) -> bytes:
        """Gera o DOCX do currículo e devolve os bytes do arquivo.

        O método prepara o documento conforme o modelo, preenche propriedades
        sem dados internos, escreve cabeçalho e seções na ordem do DTO e salva
        em memória. Ele existe para entregar à exportação um arquivo pronto
        para download.
        """
        estilo = ESTILOS_POR_MODELO.get(documento.modelo, ESTILO_CLASSICO)
        word = Document()
        _preparar_documento(word, estilo, documento.titulo_versao)
        _escrever_cabecalho(word, documento.cabecalho)
        for secao in documento.secoes:
            word.add_heading(secao.titulo, level=1)
            for bloco in secao.blocos:
                _escrever_bloco(word, bloco)
        saida = BytesIO()
        word.save(saida)
        return saida.getvalue()


def _preparar_documento(word: DocumentoWord, estilo: EstiloModelo, titulo: str) -> None:
    """Configura fonte, idioma, espaçamento, margens e propriedades do arquivo.

    A função ajusta o estilo "Normal", marca o idioma pt-BR, aplica as margens
    do modelo e substitui as propriedades padrão da biblioteca pelo título da
    versão, deixando autor e comentários vazios. Ela existe para que o arquivo
    seja legível e não carregue metadados indevidos.
    """
    normal = word.styles["Normal"]
    normal.font.name = FONTE_DOCUMENTO
    normal.font.size = Pt(estilo.tamanho_fonte)
    normal.paragraph_format.space_after = Pt(estilo.espaco_apos)
    rpr = normal.element.get_or_add_rPr()
    idioma = OxmlElement("w:lang")
    idioma.set(qn("w:val"), IDIOMA_DOCUMENTO)
    rpr.append(idioma)
    for secao in word.sections:
        secao.top_margin = secao.bottom_margin = Cm(estilo.margem_cm)
        secao.left_margin = secao.right_margin = Cm(estilo.margem_cm)
    propriedades = word.core_properties
    propriedades.title = titulo
    propriedades.language = IDIOMA_DOCUMENTO
    propriedades.author = ""
    propriedades.last_modified_by = ""
    propriedades.comments = ""
    propriedades.subject = ""
    propriedades.keywords = ""


def _escrever_cabecalho(word: DocumentoWord, cabecalho: CabecalhoDocumento) -> None:
    """Escreve o nome do estudante como título e os contatos em uma linha.

    A função usa o estilo "Title" para o nome e junta, separados por barra
    vertical, apenas os contatos informados. Ela existe para que o topo do
    arquivo corresponda à identificação mostrada na prévia.
    """
    word.add_heading(cabecalho.nome, level=0)
    contatos = [
        cabecalho.email,
        *cabecalho.telefones,
        cabecalho.cidade,
        cabecalho.endereco,
        cabecalho.linkedin,
        cabecalho.curriculo_lattes,
    ]
    word.add_paragraph(SEPARADOR_CONTATOS.join(item for item in contatos if item))


def _escrever_bloco(word: DocumentoWord, bloco: BlocoDocumento) -> None:
    """Escreve um bloco com título em negrito, subtítulo, período e detalhes.

    A função cria um parágrafo por informação presente, mantém o título junto
    do parágrafo seguinte para não separá-lo do conteúdo em quebras de página
    e escreve subtítulo e período em itálico. Ela existe para que todo tipo de
    item apareça de forma uniforme no DOCX.
    """
    titulo = word.add_paragraph()
    titulo.add_run(bloco.titulo).bold = True
    titulo.paragraph_format.keep_with_next = True
    for texto in (bloco.subtitulo, bloco.periodo):
        if texto:
            paragrafo = word.add_paragraph()
            paragrafo.add_run(texto).italic = True
            paragrafo.paragraph_format.keep_with_next = True
    for detalhe in bloco.detalhes:
        word.add_paragraph(detalhe)
