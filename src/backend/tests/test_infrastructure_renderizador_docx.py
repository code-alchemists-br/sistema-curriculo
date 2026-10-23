"""Testa a geração DOCX do currículo e seu conteúdo textual mínimo."""

import unittest
import zipfile
from dataclasses import replace
from io import BytesIO
from uuid import UUID

from docx import Document

from backend.application import (
    DocumentoCurriculoRenderizado,
    PreviaVersaoCurriculo,
    RenderizadorDocumentoCurriculo,
    montar_documento_renderizado,
)
from backend.domain import Documento, DocumentoId
from backend.infrastructure.exportacao.renderizador_docx import (
    RenderizadorDocumentoCurriculoDocx,
)
from backend.tests.fixtures.curriculo_exportacao import (
    obter_curriculo_exportacao_fixture,
)

URL_ARMAZENAMENTO = "armazenamento/interno/documento-privado.pdf"


def _documento_renderizado(com_documento: bool = False) -> DocumentoCurriculoRenderizado:
    """Monta o documento renderizado a partir da fixture estável de exportação.

    A função usa os itens da fixture como itens selecionados da versão e pode
    incluir um documento complementar com caminho sensível. Ela existe para que
    o DOCX seja gerado do mesmo DTO usado por qualquer formato.
    """
    fixture = obter_curriculo_exportacao_fixture()
    documentos = ()
    if com_documento:
        documentos = (
            Documento(
                id=DocumentoId(UUID("88888888-8888-4888-8888-888888888888")),
                usuario_id=fixture.usuario.id,
                nome_arquivo="certificado-ingles.pdf",
                tipo_arquivo="application/pdf",
                url_armazenamento=URL_ARMAZENAMENTO,
            ),
        )
    previa = PreviaVersaoCurriculo(
        usuario=fixture.usuario,
        curriculo=fixture.curriculo,
        formacoes=(fixture.formacao_academica,),
        experiencias=(fixture.experiencia_profissional,),
        projetos=(fixture.projeto_academico,),
        competencias=fixture.competencias,
        idiomas=fixture.idiomas,
        documentos=documentos,
    )
    return montar_documento_renderizado(previa)


def _paragrafos(conteudo: bytes) -> list[tuple[str, str]]:
    """Lê o DOCX gerado e devolve o texto e o estilo de cada parágrafo não vazio.

    A função abre os bytes com python-docx, como um leitor externo faria. Ela
    existe para que os testes verifiquem o arquivo produzido, e não o código
    que o produziu.
    """
    return [
        (paragrafo.text, paragrafo.style.name)
        for paragrafo in Document(BytesIO(conteudo)).paragraphs
        if paragrafo.text.strip()
    ]


class RenderizadorDocxTestCase(unittest.TestCase):
    """Verifica o arquivo DOCX gerado a partir do documento renderizado.

    Os cenários geram o arquivo em memória e o leem de volta, sem disco, banco
    ou rede. A classe existe para cobrir o conteúdo textual mínimo do UC08 e a
    correspondência com a visualização (CT-RF09-02 e CT-RF09-03).
    """

    def setUp(self) -> None:
        """Gera o DOCX padrão e lê seus parágrafos.

        O setup guarda o documento, os bytes e os parágrafos lidos. Ele existe
        para reduzir repetição entre os cenários.
        """
        self.documento = _documento_renderizado()
        self.conteudo = RenderizadorDocumentoCurriculoDocx().renderizar(self.documento)
        self.paragrafos = _paragrafos(self.conteudo)
        self.textos = [texto for texto, _ in self.paragrafos]

    def test_implementa_a_porta_e_gera_um_docx_valido(self) -> None:
        """Gera bytes de um pacote DOCX que abre normalmente.

        O teste usa o renderizador como a porta, confere a assinatura ZIP e a
        presença do corpo do Word. Ele existe para garantir que nunca seja
        entregue um arquivo inválido (CT-RF09-05).
        """
        renderizador: RenderizadorDocumentoCurriculo = RenderizadorDocumentoCurriculoDocx()

        conteudo = renderizador.renderizar(self.documento)

        self.assertTrue(conteudo.startswith(b"PK"))
        with zipfile.ZipFile(BytesIO(conteudo)) as pacote:
            self.assertIn("word/document.xml", pacote.namelist())

    def test_cabecalho_tem_nome_como_titulo_e_contatos(self) -> None:
        """Mostra o nome no estilo de título e os contatos em seguida.

        O teste confere os dois primeiros parágrafos do arquivo. Ele existe para
        que o DOCX comece pela mesma identificação da prévia.
        """
        self.assertEqual(self.paragrafos[0], ("Ana Carolina da Silva", "Title"))
        contatos = self.textos[1]
        for esperado in (
            "ana.silva@fatec.sp.gov.br",
            "(11) 98765-4321",
            "(11) 3322-1100",
            "https://linkedin.com/in/estudante-fatec",
            "http://lattes.cnpq.br/1234567890123456",
        ):
            self.assertIn(esperado, contatos)

    def test_secoes_aparecem_como_titulos_na_ordem_do_documento(self) -> None:
        """Escreve cada seção com estilo "Heading 1" e na ordem do DTO.

        O teste filtra os parágrafos de título de seção e compara com as seções
        do documento renderizado. Ele existe para que o arquivo siga a ordem da
        versão e seja navegável no Word.
        """
        titulos = [texto for texto, estilo in self.paragrafos if estilo == "Heading 1"]

        self.assertEqual(titulos, [secao.titulo for secao in self.documento.secoes])

    def test_cada_bloco_aparece_dentro_da_sua_secao(self) -> None:
        """Coloca o título de cada bloco depois do título da sua seção.

        O teste localiza, no arquivo, o título de cada seção e verifica que os
        blocos dela vêm antes da próxima seção. Ele existe para que nenhum item
        apareça na seção errada.
        """
        posicoes = [self.textos.index(secao.titulo) for secao in self.documento.secoes]
        limites = posicoes[1:] + [len(self.textos)]
        for secao, inicio, fim in zip(self.documento.secoes, posicoes, limites, strict=True):
            trecho = self.textos[inicio:fim]
            for bloco in secao.blocos:
                with self.subTest(secao=secao.titulo, bloco=bloco.titulo):
                    self.assertIn(bloco.titulo, trecho)

    def test_textos_minimos_de_periodo_e_detalhes(self) -> None:
        """Inclui períodos, subtítulos e detalhes dos itens.

        O teste procura textos característicos da fixture no arquivo. Ele
        existe para cobrir o conteúdo textual mínimo pedido pelo card.
        """
        for esperado in (
            "01/2024 - atual",
            "02/2023 - previsão 12/2025",
            "Situação: Em andamento",
            "Empresa de Soluções Tecnológicas Ltda",
            "Tecnologias: Python, FastAPI, TypeScript, React, PostgreSQL, Nix",
            "Técnica - Avançado",
            "Inglês",
        ):
            self.assertIn(esperado, self.textos)

    def test_modelos_classico_e_compacto_geram_o_mesmo_texto(self) -> None:
        """Muda apenas a aparência entre os modelos, nunca o conteúdo.

        O teste gera o DOCX nos modelos clássico e compacto e compara os textos.
        Ele existe para que trocar de modelo não altere o currículo exportado.
        """
        classico = RenderizadorDocumentoCurriculoDocx().renderizar(
            replace(self.documento, modelo="classico")
        )
        compacto = RenderizadorDocumentoCurriculoDocx().renderizar(
            replace(self.documento, modelo="compacto")
        )

        self.assertEqual(
            [texto for texto, _ in _paragrafos(classico)],
            [texto for texto, _ in _paragrafos(compacto)],
        )
        self.assertNotEqual(classico, compacto)

    def test_propriedades_e_conteudo_sem_dados_internos(self) -> None:
        """Mantém fora do arquivo IDs, hash, armazenamento e autor da biblioteca.

        O teste gera o DOCX com um documento complementar e procura dados
        internos em todo o XML do pacote, além de conferir título e idioma nas
        propriedades. Ele existe para que o arquivo exportado não revele dados
        internos nem metadados indevidos.
        """
        conteudo = RenderizadorDocumentoCurriculoDocx().renderizar(
            _documento_renderizado(com_documento=True)
        )
        fixture = obter_curriculo_exportacao_fixture()
        with zipfile.ZipFile(BytesIO(conteudo)) as pacote:
            xml = "".join(
                pacote.read(nome).decode("utf-8", "ignore")
                for nome in pacote.namelist()
                if nome.endswith(".xml")
            )
        propriedades = Document(BytesIO(conteudo)).core_properties

        self.assertIn("certificado-ingles.pdf", xml)
        for proibido in (
            URL_ARMAZENAMENTO,
            fixture.usuario.hash_senha.valor,
            str(fixture.usuario.id.valor),
            str(fixture.formacao_academica.id.valor),
        ):
            self.assertNotIn(proibido, xml)
        self.assertEqual(propriedades.title, self.documento.titulo_versao)
        self.assertEqual(propriedades.language, "pt-BR")
        self.assertEqual(propriedades.author, "")

    def test_documento_sem_secoes_gera_somente_o_cabecalho(self) -> None:
        """Gera um DOCX válido apenas com nome e contatos.

        O teste remove todas as seções do documento. Ele existe para que uma
        versão sem itens selecionados ainda produza um arquivo coerente.
        """
        conteudo = RenderizadorDocumentoCurriculoDocx().renderizar(
            replace(self.documento, secoes=())
        )

        paragrafos = _paragrafos(conteudo)
        self.assertEqual(len(paragrafos), 2)
        self.assertEqual(paragrafos[0][0], "Ana Carolina da Silva")
