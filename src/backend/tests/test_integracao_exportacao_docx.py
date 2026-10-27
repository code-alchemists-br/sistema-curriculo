# Proveniência: decision-analysis prompts/backend/20261026-082220-abertura-legibilidade-docx-v001.md#v001
"""Testes de integração para abertura, legibilidade e conformidade do DOCX gerado."""

import io
import unittest
import zipfile
from dataclasses import replace
from typing import Any
from uuid import UUID

from docx import Document

from backend.api.exportacao_curriculo import (
    CAMINHO_EXPORTACAO_DOCX,
)
from backend.api.factory import create_app
from backend.api.sessoes import COOKIE_SESSAO
from backend.application import (
    FORMATO_DOCX,
    CurriculoNaoEncontrado,
    ExportarCurriculo,
    GerarPreviaVersaoCurriculoEntrada,
    PreviaVersaoCurriculo,
)
from backend.application.autenticacao_sessao import AutenticadorSessao
from backend.domain import Documento, DocumentoId
from backend.domain.value_objects import UsuarioId
from backend.infrastructure.exportacao.renderizador_docx import (
    ESTILO_CLASSICO,
    ESTILO_COMPACTO,
    IDIOMA_DOCUMENTO,
    SEPARADOR_CONTATOS,
    RenderizadorDocumentoCurriculoDocx,
)
from backend.tests.fixtures.curriculo_exportacao import (
    ID_CURRICULO_ESTAVEL,
    ID_USUARIO_ESTAVEL,
    obter_curriculo_exportacao_fixture,
)

CREDENCIAL_AUTENTICADA = "token-integracao-docx-valido"
URL_ARMAZENAMENTO_SENSIVEL = "armazenamento/interno/privado/segredo-estudante.pdf"


class AutenticadorSessaoIntegracaoStub(AutenticadorSessao):
    """Autentica o estudante da fixture para testes integrados de exportação.

    **O que faz:** Valida a credencial sintética e devolve o identificador estável do estudante.
    **Como faz:** Compara a string de token diretamente sem acessar serviços externos ou banco de dados.
    **Qual finalidade atende:** Permite exercitar o middleware real da API de forma isolada e determinística.
    """

    async def autenticar(self, token: str) -> UsuarioId | None:
        """Devolve o UsuarioId se a credencial for válida.

        **O que faz:** Mapeia a credencial de teste para o estudante dono da versão.
        **Como faz:** Avalia a igualdade com a constante `CREDENCIAL_AUTENTICADA`.
        **Qual finalidade atende:** Autentica o contexto da requisição na camada de transporte.
        """
        return ID_USUARIO_ESTAVEL if token == CREDENCIAL_AUTENTICADA else None


class GeradorPreviaIntegracaoStub:
    """Fornece a prévia completa do currículo para o caso de uso de exportação.

    **O que faz:** Retorna a entidade de prévia populada a partir da fixture estável do projeto.
    **Como faz:** Constrói o objeto `PreviaVersaoCurriculo` contendo formações, experiências, projetos e competências.
    **Qual finalidade atende:** Habilita a execução real do caso de uso de exportação e do renderizador DOCX.
    """

    def __init__(
        self, com_documento_sensivel: bool = False, modelo: str = "classico"
    ) -> None:
        """Inicializa o stub com as opções de modelo e documentos complementares.

        **O que faz:** Configura os parâmetros de montagem da prévia.
        **Como faz:** Salva os argumentos booleanos e de layout no estado interno da instância.
        **Qual finalidade atende:** Permite alternar entre modelo clássico, compacto e documentos sensíveis.
        """
        self._com_documento_sensivel = com_documento_sensivel
        self._modelo = modelo

    async def executar(
        self, entrada: GerarPreviaVersaoCurriculoEntrada
    ) -> PreviaVersaoCurriculo:
        """Gera a prévia correspondente à entrada do caso de uso.

        **O que faz:** Valida a propriedade da versão e devolve a prévia estruturada.
        **Como faz:** Confere se o usuário e currículo coincidem com a fixture e retorna o objeto montado.
        **Qual finalidade atende:** Atua como porta de entrada de dados para o `ExportarCurriculo`.
        """
        if (entrada.usuario_id, entrada.curriculo_id) != (
            ID_USUARIO_ESTAVEL,
            ID_CURRICULO_ESTAVEL,
        ):
            raise CurriculoNaoEncontrado("Versão de currículo não encontrada.")
        fixture = obter_curriculo_exportacao_fixture()
        documentos = ()
        if self._com_documento_sensivel:
            documentos = (
                Documento(
                    id=DocumentoId(UUID("77777777-7777-4777-7777-777777777777")),
                    usuario_id=fixture.usuario.id,
                    nome_arquivo="Certificado_Pos_Graduacao.pdf",
                    tipo_arquivo="application/pdf",
                    url_armazenamento=URL_ARMAZENAMENTO_SENSIVEL,
                ),
            )
        curriculo_configurado = replace(fixture.curriculo, layout=self._modelo)
        return PreviaVersaoCurriculo(
            usuario=fixture.usuario,
            curriculo=curriculo_configurado,
            formacoes=(fixture.formacao_academica,),
            experiencias=(fixture.experiencia_profissional,),
            projetos=(fixture.projeto_academico,),
            competencias=fixture.competencias,
            idiomas=fixture.idiomas,
            documentos=documentos,
        )


def _criar_aplicacao_teste(
    com_documento_sensivel: bool = False, modelo: str = "classico"
) -> object:
    """Cria a aplicação FastAPI com os adaptadores reais e stubs autenticados.

    **O que faz:** Instancia a API completa configurando a rota de exportação DOCX.
    **Como faz:** Invoca `create_app` injetando o renderizador DOCX real e o gerador de prévia.
    **Qual finalidade atende:** Permite disparar requisições ASGI completas simulando chamadas HTTP reais.
    """
    gerador_previa = GeradorPreviaIntegracaoStub(com_documento_sensivel, modelo)
    renderizador = RenderizadorDocumentoCurriculoDocx()
    caso_de_uso = ExportarCurriculo(gerador_previa, renderizador, FORMATO_DOCX)
    return create_app(
        autenticador_sessao=AutenticadorSessaoIntegracaoStub(),
        exportar_curriculo_docx=caso_de_uso,  # type: ignore[arg-type]
    )


async def _executar_requisicao_get(
    app: object, caminho: str, autenticado: bool = True
) -> tuple[int, dict[str, str], bytes]:
    """Dispara uma requisição HTTP GET via interface ASGI sem subir servidor de rede.

    **O que faz:** Envia mensagem ASGI para a aplicação FastAPI e captura a resposta.
    **Como faz:** Constrói os dicionários de escopo e headers, despachando para a função `app`.
    **Qual finalidade atende:** Testa de ponta a ponta o contrato HTTP e os bytes retornados pelo endpoint.
    """
    headers = (
        [(b"cookie", f"{COOKIE_SESSAO}={CREDENCIAL_AUTENTICADA}".encode())]
        if autenticado
        else []
    )
    escopo = {
        "type": "http",
        "asgi": {"version": "3.0"},
        "http_version": "1.1",
        "method": "GET",
        "scheme": "http",
        "path": caminho,
        "raw_path": caminho.encode(),
        "query_string": b"",
        "headers": headers,
        "client": ("127.0.0.1", 50000),
        "server": ("127.0.0.1", 80),
    }
    status = 0
    cabecalhos_resposta: dict[str, str] = {}
    partes: list[bytes] = []

    async def receber() -> dict[str, object]:
        """Entrega o corpo vazio da requisição GET."""
        return {"type": "http.request", "body": b"", "more_body": False}

    async def enviar(mensagem: dict[str, Any]) -> None:
        """Captura os cabeçalhos e fragmentos do corpo da resposta."""
        nonlocal status
        if mensagem["type"] == "http.response.start":
            status = int(mensagem["status"])
            for nome, valor in mensagem.get("headers", []):
                cabecalhos_resposta[nome.decode("latin-1").lower()] = valor.decode(
                    "latin-1"
                )
        elif mensagem["type"] == "http.response.body":
            partes.append(mensagem.get("body", b""))

    await app(escopo, receber, enviar)  # type: ignore[operator]
    return status, cabecalhos_resposta, b"".join(partes)


class TesteIntegracaoExportacaoDocx(unittest.IsolatedAsyncioTestCase):
    """Valida a integração da exportação DOCX, sua abertura em leitor e critérios de legibilidade."""

    async def test_download_docx_abre_em_leitor_externo_e_preserva_estrutura_openxml(
        self,
    ) -> None:
        """Valida que o arquivo baixado pela rota HTTP abre perfeitamente e atende ao padrão OpenXML.

        **O que faz:** Executa a rota de download DOCX e inspeciona o pacote com leitor OpenXML e motor ZIP.
        **Como faz:** Dispara GET na rota da API, confere MIME, headers RFC 6266, partes do ZIP e lê com `docx.Document`.
        **Qual finalidade atende:** Garante que o arquivo gerado não esteja corrompido e abra em qualquer suíte de escritório (F6.14 / CT-RF09-05).
        """
        app = _criar_aplicacao_teste(com_documento_sensivel=False, modelo="classico")
        caminho = CAMINHO_EXPORTACAO_DOCX.format(
            usuario_id=ID_USUARIO_ESTAVEL.valor, curriculo_id=ID_CURRICULO_ESTAVEL.valor
        )

        status, cabecalhos, corpo = await _executar_requisicao_get(app, caminho)

        self.assertEqual(status, 200)
        self.assertEqual(cabecalhos["content-type"], FORMATO_DOCX.tipo_mime)
        self.assertIn("attachment", cabecalhos["content-disposition"])
        self.assertEqual(cabecalhos["x-content-type-options"], "nosniff")
        self.assertEqual(cabecalhos["cache-control"], "private, no-store")

        # 1. Validação da assinatura ZIP de pacotes Office Open XML (PK\x03\x04)
        self.assertTrue(corpo.startswith(b"PK\x03\x04"))

        # 2. Validação da estrutura interna do pacote OpenXML (ISO/IEC 29500)
        with zipfile.ZipFile(io.BytesIO(corpo)) as pacote:
            arquivos_pacote = pacote.namelist()
            self.assertIn("[Content_Types].xml", arquivos_pacote)
            self.assertIn("_rels/.rels", arquivos_pacote)
            self.assertIn("word/document.xml", arquivos_pacote)
            self.assertIn("word/styles.xml", arquivos_pacote)
            self.assertIn("docProps/core.xml", arquivos_pacote)

        # 3. Abertura do documento através de leitor externo (python-docx)
        doc = Document(io.BytesIO(corpo))
        self.assertIsNotNone(doc)
        self.assertGreater(len(doc.paragraphs), 5)

    async def test_legibilidade_hierarquia_e_formatacao_dos_paragrafos(self) -> None:
        """Valida os critérios tipográficos de legibilidade, hierarquia e formatação no DOCX.

        **O que faz:** Inspeciona estilos, cabeçalhos, negritos, itálicos e regras de paragrafo do documento.
        **Como faz:** Percorre os parágrafos do DOCX lido verificando níveis de título, `keep_with_next` e alinhamentos.
        **Qual finalidade atende:** Garante conformidade com RNF07 e UC08, prevenindo linhas viúvas e desorganização textual.
        """
        app = _criar_aplicacao_teste(com_documento_sensivel=False, modelo="classico")
        caminho = CAMINHO_EXPORTACAO_DOCX.format(
            usuario_id=ID_USUARIO_ESTAVEL.valor, curriculo_id=ID_CURRICULO_ESTAVEL.valor
        )

        status, _, corpo = await _executar_requisicao_get(app, caminho)
        self.assertEqual(status, 200)

        doc = Document(io.BytesIO(corpo))
        paragrafos = [p for p in doc.paragraphs if p.text.strip()]

        # O primeiro parágrafo deve conter o nome do estudante com o estilo 'Title'
        primeiro = paragrafos[0]
        self.assertEqual(primeiro.text, "Ana Carolina da Silva")
        self.assertEqual(primeiro.style.name, "Title")

        # O segundo parágrafo deve conter as informações de contato unificadas com separador
        segundo = paragrafos[1]
        self.assertIn(SEPARADOR_CONTATOS, segundo.text)
        self.assertIn("ana.silva@fatec.sp.gov.br", segundo.text)
        self.assertIn("(11) 98765-4321", segundo.text)

        # Seções devem ter estilo 'Heading 1' para navegabilidade estruturada
        titulos_secoes = [p.text for p in paragrafos if p.style.name == "Heading 1"]
        self.assertIn("Formação acadêmica", titulos_secoes)
        self.assertIn("Experiência profissional", titulos_secoes)
        self.assertIn("Projetos acadêmicos", titulos_secoes)
        self.assertIn("Competências", titulos_secoes)
        self.assertIn("Idiomas", titulos_secoes)

        # Validação de legibilidade tipográfica: títulos de blocos com keep_with_next e negrito
        for p in paragrafos:
            if p.text in (
                "Análise e Desenvolvimento de Sistemas",
                "Desenvolvedora de Software Estagiária",
                "Sistema de Apoio à Elaboração de Currículos",
            ):
                self.assertTrue(
                    p.paragraph_format.keep_with_next,
                    f"Parágrafo '{p.text}' deve ter keep_with_next ativado para evitar quebra órfã",
                )
                self.assertTrue(
                    any(run.bold for run in p.runs),
                    f"Parágrafo '{p.text}' deve possuir texto em negrito",
                )

    async def test_exportacao_modelo_compacto_mantem_mesmo_texto_com_margens_menores(
        self,
    ) -> None:
        """Valida que o Modelo Compacto preserva o mesmo conteúdo textual adotando margens e espaçamentos densos.

        **O que faz:** Gera o DOCX no modelo clássico e no modelo compacto e compara textos e margens de página.
        **Como faz:** Executa ambas as exportações, extrai os textos para atestar igualdade e afere as dimensões de margem.
        **Qual finalidade atende:** Assegura que alternar de modelo mude exclusivamente a apresentação visual e espacial sem corromper o conteúdo.
        """
        app_classico = _criar_aplicacao_teste(modelo="classico")
        app_compacto = _criar_aplicacao_teste(modelo="compacto")
        caminho = CAMINHO_EXPORTACAO_DOCX.format(
            usuario_id=ID_USUARIO_ESTAVEL.valor, curriculo_id=ID_CURRICULO_ESTAVEL.valor
        )

        _, _, corpo_classico = await _executar_requisicao_get(app_classico, caminho)
        _, _, corpo_compacto = await _executar_requisicao_get(app_compacto, caminho)

        doc_classico = Document(io.BytesIO(corpo_classico))
        doc_compacto = Document(io.BytesIO(corpo_compacto))

        textos_classico = [p.text for p in doc_classico.paragraphs if p.text.strip()]
        textos_compacto = [p.text for p in doc_compacto.paragraphs if p.text.strip()]

        # O conteúdo textual deve ser rigorosamente idêntico
        self.assertEqual(textos_classico, textos_compacto)

        # Margens do Modelo Clássico: 2.0 cm
        secao_classico = doc_classico.sections[0]
        self.assertAlmostEqual(
            secao_classico.left_margin.cm, ESTILO_CLASSICO.margem_cm, places=1
        )
        self.assertAlmostEqual(
            secao_classico.top_margin.cm, ESTILO_CLASSICO.margem_cm, places=1
        )

        # Margens do Modelo Compacto: 1.5 cm
        secao_compacto = doc_compacto.sections[0]
        self.assertAlmostEqual(
            secao_compacto.left_margin.cm, ESTILO_COMPACTO.margem_cm, places=1
        )
        self.assertAlmostEqual(
            secao_compacto.top_margin.cm, ESTILO_COMPACTO.margem_cm, places=1
        )

    async def test_higienizacao_de_metadados_e_omissao_de_urls_de_storage(self) -> None:
        """Garante que metadados internos e caminhos privados de arquivos complementares não vazem no DOCX.

        **O que faz:** Exporta currículo contendo documento anexo e inspeciona o XML bruto e propriedades do Word.
        **Como faz:** Realiza busca de strings sensíveis em todo o conteúdo descompactado e valida `core_properties`.
        **Qual finalidade atende:** Protege a privacidade do estudante e a segurança da infraestrutura de armazenamento.
        """
        app = _criar_aplicacao_teste(com_documento_sensivel=True, modelo="classico")
        caminho = CAMINHO_EXPORTACAO_DOCX.format(
            usuario_id=ID_USUARIO_ESTAVEL.valor, curriculo_id=ID_CURRICULO_ESTAVEL.valor
        )

        status, _, corpo = await _executar_requisicao_get(app, caminho)
        self.assertEqual(status, 200)

        # 1. O nome amigável do arquivo complementar deve estar presente no texto
        doc = Document(io.BytesIO(corpo))
        textos = [p.text for p in doc.paragraphs]
        self.assertTrue(any("Certificado_Pos_Graduacao.pdf" in t for t in textos))

        # 2. A URL de armazenamento privada e sensível NÃO deve constar em nenhuma parte do pacote
        with zipfile.ZipFile(io.BytesIO(corpo)) as pacote:
            xml_completo = "".join(
                pacote.read(nome).decode("utf-8", errors="replace")
                for nome in pacote.namelist()
                if nome.endswith(".xml") or nome.endswith(".rels")
            )
            self.assertNotIn(URL_ARMAZENAMENTO_SENSIVEL, xml_completo)

        # 3. Metadados do documento: idioma pt-BR configurado e ausência de dados do autor da biblioteca
        propriedades = doc.core_properties
        self.assertEqual(propriedades.language, IDIOMA_DOCUMENTO)
        self.assertEqual(propriedades.author, "")
        self.assertEqual(propriedades.last_modified_by, "")
        self.assertEqual(propriedades.comments, "")


if __name__ == "__main__":
    unittest.main()
